"""Módulo de Importação Multiformato de Modelos Conceituais e Lógicos.

Suporta importação a partir de:
1. JSON estruturado (nativo ou externo, ex.: 03-diagrama-estrela-navathe.json).
2. Planilhas Excel de Rastreabilidade DE-PARA (ex.: DEXPARA_Transacional_DW.xlsx).
3. Scripts SQL DDL (CREATE TABLE, PKs, FKs).
4. Metadados do Dicionário de Dados Oracle (via sessão ativa).
"""
import math
import os
import re
import json
import openpyxl

from models import Project, Entity, Attribute, Relationship, Domain, RelationshipArc, new_id


def auto_layout_star(entities, rels, center_x=550, center_y=420, radius=320):
    """Posiciona a entidade Fato (ou a mais conectada) no centro e as demais ao redor em círculo."""
    if not entities:
        return
    # Identifica fato pelo papel ou pelo número de relacionamentos
    fact_ent = next((e for e in entities if getattr(e, "kimball_role", "") == "fact" or e.name.upper().startswith("FT_")), None)
    if not fact_ent:
        # Pega a entidade que mais aparece nos relacionamentos
        degree = {}
        for r in rels:
            for eid, *_ in r.participants():
                degree[eid] = degree.get(eid, 0) + 1
        sorted_ents = sorted(entities, key=lambda e: degree.get(e.id, 0), reverse=True)
        fact_ent = sorted_ents[0] if sorted_ents else entities[0]

    fact_ent.x = center_x
    fact_ent.y = center_y

    others = [e for e in entities if e.id != fact_ent.id]
    n = len(others)
    if n == 0:
        return

    # Ajusta o raio conforme a quantidade de tabelas
    r = max(radius, 45 * n)
    # Define centro garantindo margem positiva
    c_x = max(center_x, int(r * 1.4) + 120)
    c_y = max(center_y, int(r) + 120)
    fact_ent.x = c_x
    fact_ent.y = c_y

    for i, e in enumerate(others):
        angle = (2 * math.pi * i) / n - (math.pi / 2)  # começa no topo
        e.x = int(c_x + r * math.cos(angle) * 1.35)  # proporção widescreen
        e.y = int(c_y + r * math.sin(angle))


def auto_layout_grid(entities, start_x=80, start_y=80, cols=4, gap_x=260, gap_y=220):
    """Distribui entidades em grade ordenada."""
    for i, e in enumerate(entities):
        row = i // cols
        col = i % cols
        e.x = start_x + col * gap_x
        e.y = start_y + row * gap_y


def import_from_json(filepath_or_dict):
    """Carrega um projeto ER a partir de um caminho de arquivo JSON ou dict."""
    if isinstance(filepath_or_dict, str):
        with open(filepath_or_dict, "r", encoding="utf-8") as f:
            data = json.load(f)
    else:
        data = filepath_or_dict

    project = Project.from_dict(data)

    # Se as entidades não tiverem posições ou estiverem todas no mesmo ponto:
    positions = {(e.x, e.y) for e in project.entities}
    if len(positions) <= 1 and len(project.entities) > 1:
        auto_layout_grid(project.entities)

    return project


def import_from_dexpara_excel(filepath):
    """Importa modelo dimensional/conceitual a partir de planilha Excel de DE-PARA.

    Lê abas padrão:
    - 2_DEPARA_Tabelas (Tabelas Fato / Dimensão, SCD, Descrição)
    - 3_DEPARA_Atributos (Campos, tipos de dados, chaves primárias e estrangeiras)
    - 4_Dominios_Valores (Domínios e valores de negócio)
    """
    wb = openpyxl.load_workbook(filepath, data_only=True)
    project = Project(notation="chen")

    # 1. Carrega Domínios (aba 4_Dominios_Valores ou similar)
    sheet_dom_name = next((s for s in wb.sheetnames if "DOMINIO" in s.upper()), None)
    domains_by_name = {}
    if sheet_dom_name:
        ws_dom = wb[sheet_dom_name]
        for row in ws_dom.iter_rows(min_row=2, values_only=True):
            if not any(row):
                continue
            # Procura coluna de campo e descrição/código
            # Tabela Origem(0), Campo Origem(1), Código(2), Descrição(3), Tabela Destino(4), Campo Destino(5)
            campo_dw = str(row[5] or row[1] or "").strip()
            cod = str(row[2] or "").strip()
            desc = str(row[3] or "").strip()
            if campo_dw and cod:
                if campo_dw not in domains_by_name:
                    d = Domain(name=f"DOM_{campo_dw}", type="STRING", description=f"Valores permitidos de {campo_dw}")
                    domains_by_name[campo_dw] = d
                    project.domains.append(d)
                d = domains_by_name[campo_dw]
                if not any(v.get("code") == cod for v in d.values):
                    d.values.append({"code": cod, "label": desc})

    # 2. Carrega Tabelas (aba 2_DEPARA_Tabelas ou similar)
    sheet_tbl_name = next((s for s in wb.sheetnames if "DEPARA_TABELA" in s.upper() or "TABELAS" in s.upper()), None)
    ent_map = {}  # nome_tabela -> Entity

    if sheet_tbl_name:
        ws_tbl = wb[sheet_tbl_name]
        for row in ws_tbl.iter_rows(min_row=2, values_only=True):
            if not any(row):
                continue
            # Colunas esperadas: Sistema Origem(0), Tabela Origem(1), Descrição(2), Tabela Destino DW(3), Tipo DW(4), SCD(5)
            tbl_dw = str(row[3] or row[1] or "").strip()
            if not tbl_dw or tbl_dw in ent_map:
                continue
            tipo_dw = str(row[4] or "").strip().lower()
            scd = str(row[5] or "").strip().lower()
            desc = str(row[2] or "").strip()

            role = "fact" if "fato" in tipo_dw or tbl_dw.startswith("FT_") else "dimension"
            scd_type = "scd2" if "2" in scd else "scd1"

            ent = Entity(name=tbl_dw, is_weak=False, kimball_role=role, scd_type=scd_type, description=desc, attrs=[])
            ent_map[tbl_dw] = ent
            project.add_entity(ent)

    # 3. Carrega Atributos (aba 3_DEPARA_Atributos ou similar)
    sheet_att_name = next((s for s in wb.sheetnames if "DEPARA_ATRIBUTO" in s.upper() or "ATRIBUTOS" in s.upper()), None)
    fk_links = []  # (tabela_fato, tabela_dim, nome_campo)

    if sheet_att_name:
        ws_att = wb[sheet_att_name]
        for row in ws_att.iter_rows(min_row=2, values_only=True):
            if not any(row):
                continue
            # Colunas: ID(0), Tabela Origem(1), Campo Origem(2), Tipo Origem(3), Tabela Destino DW(4), Campo Destino DW(5), Tipo DW(6), Papel(7)
            tbl_dw = str(row[4] or row[1] or "").strip()
            campo_dw = str(row[5] or row[2] or "").strip()
            if not tbl_dw or not campo_dw:
                continue

            if tbl_dw not in ent_map:
                role = "fact" if tbl_dw.startswith("FT_") else "dimension"
                ent = Entity(name=tbl_dw, kimball_role=role, attrs=[])
                ent_map[tbl_dw] = ent
                project.add_entity(ent)

            ent = ent_map[tbl_dw]

            # Detecta tipo
            tipo_raw = str(row[6] or row[3] or "STRING").upper()
            if "INT" in tipo_raw or "NUMBER" in tipo_raw:
                tipo = "INTEGER"
            elif "DATE" in tipo_raw or "TIME" in tipo_raw:
                tipo = "DATE"
            elif "DECIMAL" in tipo_raw or "NUMERIC" in tipo_raw or "FLOAT" in tipo_raw:
                tipo = "DECIMAL"
            elif "BOOL" in tipo_raw:
                tipo = "BOOLEAN"
            else:
                tipo = "STRING"

            # Detecta PK
            is_pk = False
            seq = row[0]
            if seq in (1, "1") or campo_dw.startswith("ID_") and (campo_dw == f"ID_{tbl_dw.replace('FT_', '').replace('DM_', '')}"):
                if not any(a.is_pk for a in ent.attrs):
                    is_pk = True

            # Detecta FK na Fato
            if ent.kimball_role == "fact" and campo_dw.startswith("ID_"):
                # Procura a dimensão correspondente
                dim_candidate = next((dname for dname in ent_map if dname != tbl_dw and dname.replace("DM_", "") in campo_dw), None)
                if dim_candidate:
                    fk_links.append((tbl_dw, dim_candidate, campo_dw))

            # Verifica se tem domínio associado
            dom = domains_by_name.get(campo_dw)

            # Evita atributo duplicado
            if not any(a.name == campo_dw for a in ent.attrs):
                attr = Attribute(name=campo_dw, type=tipo, pk=is_pk, nn=is_pk,
                                 domain_id=dom.id if dom else None,
                                 description=str(row[8] or row[2] or ""))
                ent.attrs.append(attr)

    # 4. Cria os Relacionamentos Fato <-> Dimensão
    created_pairs = set()
    fact_entities = [e for e in project.entities if e.kimball_role == "fact" or e.name.startswith("FT_")]
    dim_entities = [e for e in project.entities if e.kimball_role != "fact" and not e.name.startswith("FT_")]

    if fact_entities:
        fact = fact_entities[0]
        # Conecta todas as dimensões à Fato
        for dim in dim_entities:
            pair = tuple(sorted([fact.id, dim.id]))
            if pair in created_pairs:
                continue
            created_pairs.add(pair)
            rel_name = f"REL_{dim.name.replace('DM_', '')[:18]}"
            rel = Relationship(
                entity1_id=dim.id,
                entity2_id=fact.id,
                card1="1",
                card2="N",
                name=rel_name,
                part1="total",
                part2="partial",
                description=f"Relacionamento dimensional entre {dim.name} e {fact.name}",
            )
            project.rels.append(rel)

    # 5. Aplica layout em estrela automático
    auto_layout_star(project.entities, project.rels)

    return project


def import_from_sql_ddl(ddl_text):
    """Importa modelo a partir de instruções SQL DDL (CREATE TABLE, PRIMARY KEY, FOREIGN KEY)."""
    project = Project(notation="chen")
    table_pattern = re.compile(r"CREATE\s+TABLE\s+(?:IF\s+NOT\s+EXISTS\s+)?([A-Za-z0-9_.\"']+)\s*\((.*?)\);", re.IGNORECASE | re.DOTALL)
    
    fk_statements = []
    entities_by_name = {}

    for match in table_pattern.finditer(ddl_text):
        raw_tbl_name = match.group(1).replace('"', '').replace("'", '').split('.')[-1].strip()
        body = match.group(2)

        role = "fact" if raw_tbl_name.upper().startswith("FT_") else "dimension" if raw_tbl_name.upper().startswith("DM_") else "entity"
        ent = Entity(name=raw_tbl_name, is_weak=False, kimball_role=role)
        entities_by_name[raw_tbl_name.upper()] = ent
        project.add_entity(ent)

        lines = [line.strip().rstrip(',') for line in body.split('\n') if line.strip()]
        for line in lines:
            if re.match(r"^(?:CONSTRAINT\s+\w+\s+)?PRIMARY\s+KEY", line, re.IGNORECASE):
                pk_cols = re.findall(r"\b([A-Za-z0-9_]+)\b", line[line.upper().find("KEY"):])
                for col in pk_cols:
                    col_attr = next((a for a in ent.attrs if a.name.upper() == col.upper()), None)
                    if col_attr:
                        col_attr.is_primary = True
                continue

            if re.match(r"^(?:CONSTRAINT\s+\w+\s+)?FOREIGN\s+KEY", line, re.IGNORECASE):
                fk_statements.append((raw_tbl_name, line))
                continue

            # Coluna normal: nome tipo [NOT NULL] [PRIMARY KEY]
            parts = line.split()
            if len(parts) >= 2 and not parts[0].upper() in ("CONSTRAINT", "CHECK", "UNIQUE", "KEY", "INDEX"):
                cname = parts[0].replace('"', '').replace("'", '').strip()
                ctype_raw = parts[1].upper()
                tipo = "STRING"
                if "INT" in ctype_raw or "NUMBER" in ctype_raw:
                    tipo = "INTEGER"
                elif "DATE" in ctype_raw or "TIME" in ctype_raw:
                    tipo = "DATE"
                elif "DECIMAL" in ctype_raw or "NUMERIC" in ctype_raw or "FLOAT" in ctype_raw:
                    tipo = "DECIMAL"
                elif "BOOL" in ctype_raw:
                    tipo = "BOOLEAN"

                is_pk = "PRIMARY" in line.upper() and "KEY" in line.upper()
                is_opt = "NOT NULL" not in line.upper() and not is_pk

                attr = Attribute(name=cname, type=tipo, is_primary=is_pk, is_optional=is_opt)
                ent.attrs.append(attr)

    # Resolve FKs
    for child_tbl, fk_line in fk_statements:
        ref_match = re.search(r"REFERENCES\s+([A-Za-z0-9_.\"']+)", fk_line, re.IGNORECASE)
        if ref_match:
            parent_tbl = ref_match.group(1).replace('"', '').replace("'", '').split('.')[-1].strip()
            c_ent = entities_by_name.get(child_tbl.upper())
            p_ent = entities_by_name.get(parent_tbl.upper())
            if c_ent and p_ent:
                rel_name = f"FK_{child_tbl}_{parent_tbl}"[:28]
                rel = Relationship(
                    entity1_id=p_ent.id,
                    entity2_id=c_ent.id,
                    card1="1",
                    card2="N",
                    name=rel_name,
                    part1="partial",
                    part2="total",
                )
                project.rels.append(rel)

    # Layout automático
    if any(e.kimball_role == "fact" for e in project.entities):
        auto_layout_star(project.entities, project.rels)
    else:
        auto_layout_grid(project.entities)

    return project


def import_from_oracle_tables(sess, schema, table_names):
    """Importa tabelas e relacionamentos FK diretamente do dicionário de dados Oracle via sessão ativa."""
    import oracle_tools as O

    project = Project(notation="chen")
    ent_map = {}
    fk_list = []

    for tname in table_names:
        det = O.table_details(sess.conn, sess.views, schema, tname)
        role = "fact" if tname.upper().startswith("FT_") else ("dimension" if tname.upper().startswith("DM_") else "entity")
        summary = det.get("summary", {})
        ent = Entity(name=tname, kimball_role=role, description=summary.get("comment", ""), attrs=[])

        pk_cols = set(det.get("pk", []))
        for col in det.get("columns", []):
            cname = col.get("name", "")
            ctype_raw = col.get("type", "VARCHAR2").upper()
            tipo = "STRING"
            if "NUMBER" in ctype_raw or "INT" in ctype_raw:
                tipo = "INTEGER" if ("(0)" in ctype_raw or "," not in ctype_raw) else "DECIMAL"
            elif "DATE" in ctype_raw or "TIME" in ctype_raw:
                tipo = "DATE"
            elif "CLOB" in ctype_raw or "BLOB" in ctype_raw:
                tipo = "TEXT"

            is_pk = cname in pk_cols
            is_nn = not col.get("nullable", True)
            attr = Attribute(name=cname, type=tipo, pk=is_pk, nn=is_nn or is_pk,
                             description=col.get("comment", ""))
            ent.attrs.append(attr)

        ent_map[tname.upper()] = ent
        project.add_entity(ent)

        for cons in det.get("constraints", []):
            if cons.get("type") == "R" and cons.get("ref_table"):
                fk_list.append((tname.upper(), cons["ref_table"].upper(), cons.get("name", "")))

    # Cria relacionamentos para as FKs entre as tabelas importadas
    for child_name, parent_name, cons_name in fk_list:
        if child_name in ent_map and parent_name in ent_map:
            c_ent = ent_map[child_name]
            p_ent = ent_map[parent_name]
            rel_name = cons_name or f"FK_{child_name}_{parent_name}"[:28]
            rel = Relationship(
                entity1_id=p_ent.id,
                entity2_id=c_ent.id,
                card1="1",
                card2="N",
                name=rel_name,
                part1="partial",
                part2="total",
                description=f"Chave estrangeira de {child_name} para {parent_name}",
            )
            project.rels.append(rel)

    if any(e.kimball_role == "fact" for e in project.entities):
        auto_layout_star(project.entities, project.rels)
    else:
        auto_layout_grid(project.entities)

    return project

